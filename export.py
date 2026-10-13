"""Exportación: CSV corregido, libro Excel (importable a Google Sheets) y ZIP de entregables."""
from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .changelog import ISSUE_COLS, LOG_COLS
from .pipeline import CleanResult, consolidate_facts

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
HEADER_FILL = PatternFill("solid", fgColor="1F3A2E")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def slug(text: str, maxlen: int = 60) -> str:
    s = re.sub(r"[^A-Za-z0-9]+", "_", text.encode("ascii", "ignore").decode() or text).strip("_")
    return (s or "tabla")[:maxlen]


def to_csv_bytes(df: pd.DataFrame, sep: str = ",", decimal: str = ".") -> bytes:
    out = df.copy()
    for c in out.columns:
        if "datetime" in str(out[c].dtype):
            out[c] = out[c].dt.strftime("%Y-%m-%d")
    return out.to_csv(index=False, sep=sep, decimal=decimal, float_format="%.10g").encode("utf-8-sig")


def summary_frame(results: list[CleanResult]) -> pd.DataFrame:
    rows = []
    for r in results:
        rows.append(
            {
                "tabla": r.table_id, "fuente": r.source, "hoja": r.sheet, "titulo_detectado": r.titulo,
                "filas_entrada": r.shape_in[0], "cols_entrada": r.shape_in[1],
                "filas_salida": r.shape_out[0], "cols_salida": r.shape_out[1],
                "celdas_corregidas": r.n_changes, "anomalias": 0 if r.issues.empty else len(r.issues),
                "calidad_antes": r.quality_before.get("puntaje"), "calidad_despues": r.quality_after.get("puntaje"),
                "pii_excluido": ", ".join(r.pii) if r.pii else "", "en_base_consolidada": "sí" if (r.facts is not None and not r.facts.empty) else "no",
                "estado": r.skipped or "OK",
            }
        )
    return pd.DataFrame(rows)


def quality_frame(results: list[CleanResult]) -> pd.DataFrame:
    rows = []
    for r in results:
        if not r.quality_before:
            continue
        for fase, q in (("antes", r.quality_before), ("despues", r.quality_after)):
            rows.append({"tabla": r.table_id, "fase": fase, **q})
    return pd.DataFrame(rows)


def all_logs(results: list[CleanResult]) -> pd.DataFrame:
    frames = [r.log for r in results if not r.log.empty]
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=LOG_COLS)


def all_issues(results: list[CleanResult]) -> pd.DataFrame:
    frames = [r.issues for r in results if not r.issues.empty]
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=ISSUE_COLS)


def all_dictionary(results: list[CleanResult]) -> pd.DataFrame:
    frames = []
    for r in results:
        if not r.dictionary.empty:
            d = r.dictionary.copy()
            d.insert(0, "tabla", r.table_id)
            frames.append(d)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def _sanitize(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for c in out.columns:
        if out[c].dtype == object or str(out[c].dtype) in ("str", "string"):
            out[c] = out[c].map(lambda v: _ILLEGAL.sub("", v) if isinstance(v, str) else v)
    return out


def _style(ws, df: pd.DataFrame):
    for j, col in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=j)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sample = df[col].head(200).map(lambda v: len(str(v)))
        width = max([len(str(col))] + sample.tolist()) if len(sample) else len(str(col))
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, width + 2), 48)
        if "datetime" in str(df[col].dtype):
            for row in range(2, len(df) + 2):
                ws.cell(row=row, column=j).number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    if len(df.columns):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{max(len(df) + 1, 2)}"


def _sheet_name(base: str, used: set[str]) -> str:
    name = re.sub(r"[\[\]\*\?/\\:]", "_", base)[:31]
    k, cand = 1, name
    while cand in used:
        k += 1
        cand = f"{name[:28]}_{k}"
    used.add(cand)
    return cand


def build_workbook(results: list[CleanResult], title: str = "Base_Reportes_Operativos_LaArena") -> bytes:
    facts = consolidate_facts(results)
    summ = summary_frame(results)
    ok = [r for r in results if not r.df.empty]
    leeme = pd.DataFrame(
        {
            "Campo": ["Documento", "Generado", "Herramienta", "Tablas procesadas", "Tablas con datos", "Registros en base consolidada",
                      "Principio", "Trazabilidad", "Datos personales", "Cómo usar en Google Sheets"],
            "Detalle": [
                title, datetime.now().strftime("%Y-%m-%d %H:%M"), "Digitalizador La Arena (arena_core)", len(results), len(ok), len(facts),
                "No se inventan datos: los vacíos se mantienen vacíos y los valores imposibles se anulan con registro.",
                "Hoja BITACORA: cada corrección con regla, columna, nº de celdas, ejemplo antes/después. Hoja ANOMALIAS: celdas puntuales.",
                "DNI, nombres, direcciones, teléfonos y fechas de nacimiento se excluyen o se pseudonimizan (ver INDICE/pii_excluido).",
                "Google Drive → Nuevo → Subir archivo → abrir con Hojas de cálculo (o Archivo → Importar en un Sheet vacío).",
            ],
        }
    )
    used: set[str] = set()
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="yyyy-mm-dd") as xw:
        fixed = [("LEEME", leeme), ("INDICE", summ), ("CALIDAD", quality_frame(results)), ("HECHOS_CONSOLIDADO", facts),
                 ("BITACORA", all_logs(results)), ("ANOMALIAS", all_issues(results)), ("DICCIONARIO", all_dictionary(results))]
        for name, df in fixed:
            df = _sanitize(df)
            df.to_excel(xw, sheet_name=_sheet_name(name, used), index=False)
            _style(xw.sheets[list(xw.sheets)[-1]], df)
        for i, r in enumerate(ok, 1):
            df = _sanitize(r.df)
            sn = _sheet_name(f"T{i:02d}_{slug(r.sheet or r.source, 24)}", used)
            df.to_excel(xw, sheet_name=sn, index=False)
            _style(xw.sheets[sn], df)
    return buf.getvalue()


def build_zip(results: list[CleanResult], sep: str = ",", decimal: str = ".") -> bytes:
    facts = consolidate_facts(results)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for i, r in enumerate([r for r in results if not r.df.empty], 1):
            z.writestr(f"tablas_corregidas/{i:02d}_{slug(r.table_id)}.csv", to_csv_bytes(r.df, sep, decimal))
            if r.long is not None:
                z.writestr(f"tablas_formato_largo/{i:02d}_{slug(r.table_id)}_largo.csv", to_csv_bytes(r.long, sep, decimal))
        z.writestr("consolidado/hechos_operativos.csv", to_csv_bytes(facts, sep, decimal))
        z.writestr("reportes/bitacora_correcciones.csv", to_csv_bytes(all_logs(results), sep, decimal))
        z.writestr("reportes/anomalias.csv", to_csv_bytes(all_issues(results), sep, decimal))
        z.writestr("reportes/diccionario_datos.csv", to_csv_bytes(all_dictionary(results), sep, decimal))
        z.writestr("reportes/indice_tablas.csv", to_csv_bytes(summary_frame(results), sep, decimal))
        z.writestr("Base_Reportes_Operativos_LaArena.xlsx", build_workbook(results))
        z.writestr("LEEME.txt", "Entregables del Digitalizador La Arena.\n- tablas_corregidas/: un CSV limpio por hoja/tabla.\n- consolidado/hechos_operativos.csv: base única en formato largo (fecha, dimensión, indicador, unidad, valor).\n- reportes/: bitácora de correcciones, anomalías, diccionario y calidad.\n- Base_Reportes_Operativos_LaArena.xlsx: libro listo para subir a Google Sheets.\nNo se inventaron datos; todo cambio está en la bitácora.\n")
    return buf.getvalue()
